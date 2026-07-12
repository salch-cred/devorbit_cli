import base64, json, os, subprocess, tempfile, time, unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import Mock, patch

from acli.extended import checkpoints, documents, plugins, project
from acli.extended.mcp_client import MCPClient
from acli.extended.router import ResponseCache
from acli.loop_engine import LoopEngine
from acli.production.native_sandbox import run_restricted
from acli.tools import browser_tools, fs_tools, git_tools, github_tools, web_tools
from acli.tools.policy import PermissionDenied, ensure_network_allowed


def settings(**kw):
    data=dict(allow_file_reads=True,allow_file_writes=True,denied_file_patterns=['.env','*.pem'],network_default='allow',allowed_domains=[],denied_domains=[],browser_allowed_domains=[],browser_denied_domains=[],browser_javascript_policy='ask',workspace_dir='.',github_token='t',github_repo='o/r',system_prompt='s',max_context_tokens=20,max_retries_per_model=1,backoff_base_seconds=0,temperature=.2,verbose_agent_chat=False,max_tool_iterations=2,browser_confirm_actuation=True,autosave_conversations=False,history_dir='.',max_saved_conversations=2,save_tool_results=True,notifications=False,response_cache=False,auto_route=False,terminal_policy='isolated_only',sandbox_backend='auto',allow_native_terminal=False,mcp_tools_enabled=True,provider='nvidia')
    data.update(kw); return SimpleNamespace(**data)


class Resp:
    def __init__(self,status=200,data=None,text='',ctype='application/json'):
        self.status_code=status; self._data=data or {}; self.text=text; self.url='https://example.com'; self.headers={'content-type':ctype}
    def json(self): return self._data
    def raise_for_status(self):
        if self.status_code>=400: raise RuntimeError(str(self.status_code))


class ComprehensiveTests(unittest.TestCase):
    def test_file_policy_and_checkpoint_secrets(self):
        with tempfile.TemporaryDirectory() as d:
            s=settings(); fs_tools.write_file(d,'a.txt','x',settings=s)
            self.assertEqual(fs_tools.read_file(d,'a.txt',settings=s),'x')
            with self.assertRaises(Exception): fs_tools.write_file(d,'.env','secret',settings=s)
            Path(d,'.env').write_text('TOKEN=x'); Path(d,'cert.pem').write_text('x')
            name=checkpoints.create_checkpoint(d,'safe').split()[2]
            root=Path(d,'.devorbit-checkpoints',name)
            self.assertFalse((root/'.env').exists()); self.assertFalse((root/'cert.pem').exists())

    def test_network_policy(self):
        s=settings(network_default='deny',allowed_domains=['example.com'],denied_domains=['bad.example.com'])
        ensure_network_allowed('https://api.example.com',s)
        with self.assertRaises(PermissionDenied): ensure_network_allowed('https://other.test',s)
        with self.assertRaises(PermissionDenied): ensure_network_allowed('https://bad.example.com',s)

    def test_native_runner_scrubs_secrets(self):
        with tempfile.TemporaryDirectory() as d, patch.dict(os.environ,{'AWS_SECRET_ACCESS_KEY':'leak','CUSTOM_TOKEN':'leak'}):
            Path(d,'check_env.py').write_text("import os\nprint(os.getenv('AWS_SECRET_ACCESS_KEY'))\n")
            out=json.loads(run_restricted(d,'python check_env.py'))
            self.assertNotIn('leak',out['stdout'])

    def test_history_trim_turn_boundary(self):
        with tempfile.TemporaryDirectory() as d:
            e=LoopEngine(Mock(),['m'],settings(workspace_dir=d,history_dir=d),enable_tools=False)
            e.messages=[{'role':'system','content':'s'},{'role':'user','content':'old'*20},{'role':'assistant','content':'answer'*20},{'role':'user','content':'new'},{'role':'assistant','content':'ok'}]
            e._trim_history(); self.assertEqual([m['role'] for m in e.messages],['system','user','assistant'])

    def test_response_cache_closes_connections(self):
        with tempfile.TemporaryDirectory() as d:
            cache=ResponseCache(Path(d,'c.sqlite')); key=cache.key('m',[{'role':'user','content':'x'}]); cache.put(key,{'content':'y'}); self.assertEqual(cache.get(key)['content'],'y')

    def test_mcp_stdio_timeout_and_response(self):
        with tempfile.TemporaryDirectory() as d:
            server=Path(d,'s.py'); server.write_text("import sys,json\nfor line in sys.stdin:\n m=json.loads(line)\n if m.get('id')==2: print(json.dumps({'jsonrpc':'2.0','id':2,'result':{'tools':[]}}),flush=True)\n")
            self.assertEqual(MCPClient(f'python {server}')._exchange('tools/list',timeout=1),{'tools':[]})
            slow=Path(d,'slow.py'); slow.write_text('import time; time.sleep(2)')
            start=time.time()
            with self.assertRaises(TimeoutError): MCPClient(f'python {slow}')._exchange('tools/list',timeout=.2)
            self.assertLess(time.time()-start,1)

    def test_browser_download_boundary(self):
        with tempfile.TemporaryDirectory() as d, tempfile.TemporaryDirectory() as outside:
            with self.assertRaises(browser_tools.ToolError): browser_tools.BrowserSession().configure(True,outside,settings(workspace_dir=d))
            session=browser_tools.BrowserSession(); session.configure(True,str(Path(d,'downloads')),settings(workspace_dir=d)); self.assertTrue(Path(d,'downloads').exists())

    def test_github_and_web_mocks(self):
        s=settings()
        with patch('requests.post',return_value=Resp(data={'number':1,'html_url':'u'})): self.assertIn('#1',github_tools.create_issue(s,'x'))
        encoded=base64.b64encode(b'hello').decode()
        with patch('requests.get',return_value=Resp(data={'content':encoded})): self.assertEqual(github_tools.get_file(s,'a'),'hello')
        html='<div class="result"><a class="result__a" href="https://example.com">R</a><div class="result__snippet">S</div></div>'
        with patch('requests.get',return_value=Resp(text=html,ctype='text/html')): self.assertIn('R',web_tools.web_search('q'))

    def test_documents_plugins_git_index(self):
        from docx import Document
        from openpyxl import Workbook
        with tempfile.TemporaryDirectory() as d:
            Path(d,'a.txt').write_text('plain'); doc=Document(); doc.add_paragraph('doc'); doc.save(Path(d,'a.docx')); wb=Workbook(); wb.active['A1']='sheet'; wb.save(Path(d,'a.xlsx'))
            self.assertIn('plain',documents.read_document(d,'a.txt')); self.assertIn('doc',documents.read_document(d,'a.docx')); self.assertIn('sheet',documents.read_document(d,'a.xlsx'))
            subprocess.run(['git','init','-q'],cwd=d,check=True); subprocess.run(['git','config','user.email','t@e'],cwd=d); subprocess.run(['git','config','user.name','T'],cwd=d)
            self.assertIn('a.txt',git_tools.git_status(d)); self.assertIn('Indexed',project.build_index(d))
        root=str(Path(__file__).resolve().parent.parent); self.assertIn('"ok": true',plugins.run_plugin(root,'example',{}).lower())


if __name__=='__main__': unittest.main()
